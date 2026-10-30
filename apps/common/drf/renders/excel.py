from io import BytesIO
import json
import re
from tempfile import TemporaryFile

from .base import BaseFileRenderer

ILLEGAL_CHARACTERS_RE = re.compile(r'[\000-\010]|[\013-\014]|[\016-\037]')


class ExcelFileRenderer(BaseFileRenderer):
    media_type = "application/xlsx"
    format = "xlsx"

    wb = None
    ws = None
    row_count = 0

    def initial_writer(self):
        from openpyxl import Workbook

        self.row_count = 0
        self.column_lengths = []
        self._row_file = TemporaryFile(mode='w+t', encoding='utf-8')
        self.wb = Workbook(write_only=True)
        self.ws = self.wb.create_sheet()

    def write_row(self, row):
        self.row_count += 1
        values = [ILLEGAL_CHARACTERS_RE.sub('', str(value)) for value in row]
        for index, value in enumerate(values):
            if index == len(self.column_lengths):
                self.column_lengths.append(0)
            self.column_lengths[index] = max(self.column_lengths[index], len(value))
        # Column dimensions precede sheet data in write-only XLSX. Spool rows
        # privately to disk so widths can still reflect every exported value.
        self._row_file.write(json.dumps(values, ensure_ascii=False) + '\n')

    def after_render(self):
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.utils import get_column_letter

        for index, length in enumerate(self.column_lengths, start=1):
            self.ws.column_dimensions[get_column_letter(index)].width = min(300, max(30, length + 2))
        self._row_file.seek(0)
        for index, line in enumerate(self._row_file, start=1):
            self.ws.row_dimensions[index].height = 20
            cells = []
            for value in json.loads(line):
                cell = WriteOnlyCell(self.ws, value=value)
                cell.data_type = 's'
                cells.append(cell)
            self.ws.append(cells)
            # The row XML is already written; do not accumulate row dimensions.
            del self.ws.row_dimensions[index]

    def get_rendered_value(self):
        buffer = BytesIO()
        self.wb.save(buffer)
        return buffer.getvalue()

    def close_writer(self):
        row_file = getattr(self, '_row_file', None)
        if row_file is not None:
            row_file.close()
            self._row_file = None
